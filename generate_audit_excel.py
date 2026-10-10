import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Create workbook and sheet
wb = openpyxl.Workbook()

# Sheet 1: Executive Summary & System Status
ws1 = wb.active
ws1.title = "Project Audit Summary"
ws1.views.sheetView[0].showGridLines = True

# Title Block
ws1.merge_cells("A1:F1")
title_cell = ws1["A1"]
title_cell.value = "GitHub Project Monitoring Agent - System Comprehensive Audit"
title_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
title_cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
title_cell.alignment = Alignment(horizontal="center", vertical="center")
ws1.row_dimensions[1].height = 40

# Section Header
ws1.merge_cells("A3:F3")
sec1 = ws1["A3"]
sec1.value = "1. Subsystem Health Overview"
sec1.font = Font(name="Calibri", size=13, bold=True, color="1F4E78")
ws1.row_dimensions[3].height = 25

headers1 = ["Module", "Component", "Current Status", "Key Issues Identified", "Severity", "Action Required"]
for col_idx, h in enumerate(headers1, 1):
    cell = ws1.cell(row=4, column=col_idx, value=h)
    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws1.row_dimensions[4].height = 28

audit_data_summary = [
    (
        "FastAPI-AI-Services",
        "LLM Providers (gemini_provider.py)",
        "RESOLVED",
        "SyntaxError: '(' was never closed on line 14 import statement + missing logger & LLMError imports.",
        "HIGH (Blocker)",
        "FIXED: Corrected import syntax, imported logger & LLMError. main.py imports cleanly now."
    ),
    (
        "FastAPI-AI-Services",
        "Express Client Integration",
        "WARNING",
        "[ExpressClient] Connection error on GET /repositories when GitHub-Backend is offline.",
        "MEDIUM",
        "Ensure GitHub-Backend is running on port 5000 before initiating repository context syncs."
    ),
    (
        "FastAPI-AI-Services",
        "LLM Provider Rate Limits",
        "WARNING",
        "Betopia AI / External LLM API returned HTTP 429 Rate Limit Exceeded during test suite run.",
        "MEDIUM",
        "Verify LLM API key quotas or switch default AI provider to Gemini/OpenAI in .env settings."
    ),
    (
        "GitHub-Backend",
        "TypeScript Build & Compilation",
        "PASSING",
        "npx tsc --noEmit executed with 0 type errors across all backend controllers and models.",
        "LOW (Info)",
        "No type errors found. Ready for runtime execution."
    ),
    (
        "GitHub-Backend",
        "Database & Redis Dependency",
        "DEPENDENCY",
        "Backend requires running PostgreSQL (Neon/Local) and Redis for queues (BullMQ).",
        "HIGH (Runtime)",
        "Verify DATABASE_URL and REDIS_URL environment variables in .env before running npm run dev."
    ),
    (
        "GitHub-Frontend",
        "TypeScript Build & UI Components",
        "PASSING",
        "npx tsc --noEmit executed with 0 type errors across Next.js components and providers.",
        "LOW (Info)",
        "No compilation errors found. Ready for local dev server (npm run dev)."
    ),
    (
        "GitHub-Frontend",
        "API Integration & Proxy",
        "CONFIG",
        "Requires NEXT_PUBLIC_API_URL (port 5000) and NEXT_PUBLIC_AI_SERVICE_URL (port 8000).",
        "MEDIUM",
        "Ensure both Backend (5000) and AI-Services (8000) are started simultaneously."
    ),
]

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

status_fills = {
    "RESOLVED": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
    "PASSING": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
    "WARNING": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    "CONFIG": PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),
    "DEPENDENCY": PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
}

status_fonts = {
    "RESOLVED": Font(name="Calibri", size=10, bold=True, color="375623"),
    "PASSING": Font(name="Calibri", size=10, bold=True, color="375623"),
    "WARNING": Font(name="Calibri", size=10, bold=True, color="806000"),
    "CONFIG": Font(name="Calibri", size=10, bold=True, color="1B365D"),
    "DEPENDENCY": Font(name="Calibri", size=10, bold=True, color="C65911")
}

for row_idx, data in enumerate(audit_data_summary, 5):
    for col_idx, val in enumerate(data, 1):
        cell = ws1.cell(row=row_idx, column=col_idx, value=val)
        cell.font = Font(name="Calibri", size=10)
        cell.border = thin_border
        cell.alignment = Alignment(vertical="center", wrap_text=True)

        if col_idx == 3: # Status column
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if val in status_fills:
                cell.fill = status_fills[val]
                cell.font = status_fonts[val]

    ws1.row_dimensions[row_idx].height = 36

# Sheet 2: Detailed File Audit
ws2 = wb.create_sheet(title="Detailed File Audit")
ws2.views.sheetView[0].showGridLines = True

ws2.merge_cells("A1:E1")
t2 = ws2["A1"]
t2.value = "Detailed Codebase Audit Log"
t2.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
t2.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
t2.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[1].height = 40

headers2 = ["Module", "File Path", "Audit Type", "Result", "Notes & Status"]
for col_idx, h in enumerate(headers2, 1):
    cell = ws2.cell(row=3, column=col_idx, value=h)
    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[3].height = 25

detailed_files = [
    ("FastAPI-AI-Services", "app/engineering_agent/llm/gemini_provider.py", "Syntax & Import Check", "FIXED", "Unclosed parenthesis on import fixed; logger & LLMError added."),
    ("FastAPI-AI-Services", "main.py", "Application Entrypoint", "PASSING", "Imports cleanly; ready to start server with python main.py."),
    ("FastAPI-AI-Services", "app/engineering_agent/llm/factory.py", "LLM Provider Registry", "PASSING", "All providers registered correctly in AgentLLMProviderFactory."),
    ("FastAPI-AI-Services", "app/engineering_agent/llm/openai_compatible.py", "OpenAI / Betopia Provider", "PASSING", "Valid provider configuration."),
    ("GitHub-Backend", "src/server.ts", "Express Server Entrypoint", "PASSING", "Typechecks clean via tsc."),
    ("GitHub-Backend", "database/migrations/*.sql", "Database Schemas", "PASSING", "SQL migration scripts verified."),
    ("GitHub-Frontend", "src/providers/layout-provider.tsx", "Layout Context Provider", "PASSING", "Typechecks clean via tsc."),
    ("GitHub-Frontend", "src/components/projects/CreateProjectModal.tsx", "React Component", "PASSING", "Typechecks clean via tsc."),
]

for row_idx, data in enumerate(detailed_files, 4):
    for col_idx, val in enumerate(data, 1):
        cell = ws2.cell(row=row_idx, column=col_idx, value=val)
        cell.font = Font(name="Calibri", size=10)
        cell.border = thin_border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if col_idx == 4:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if val in status_fills:
                cell.fill = status_fills[val]
                cell.font = status_fonts[val]
    ws2.row_dimensions[row_idx].height = 26

# Adjust column widths for both sheets
col_widths_ws1 = [22, 32, 16, 45, 16, 45]
for col_idx, width in enumerate(col_widths_ws1, 1):
    col_letter = get_column_letter(col_idx)
    ws1.column_dimensions[col_letter].width = width

col_widths_ws2 = [22, 50, 22, 16, 50]
for col_idx, width in enumerate(col_widths_ws2, 1):
    col_letter = get_column_letter(col_idx)
    ws2.column_dimensions[col_letter].width = width

# Save file
excel_path = r"c:\Nehal devs\Nehal-Personal\Project\GitHub-Project-Monitoring-Agent\PROJECT_AUDIT_REPORT.xlsx"
wb.save(excel_path)
print(f"Excel audit report generated successfully at: {excel_path}")
